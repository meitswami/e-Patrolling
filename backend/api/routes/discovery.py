"""Discovery API — auto-discovered targets list, export, and manual trigger."""

import asyncio
import csv
import io
import json
import math
from datetime import date, datetime, timezone
from typing import Optional
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from auth.middleware import get_current_user
from auth.rbac import require_admin, require_analyst
from database.connection import get_session
from database.models import User
from database.models_discovery import DiscoveredTarget, DiscoveryRun
from services.discovery_crawler import convert_discovered_to_target, run_discovery
from services.discovery_state import discovery_state

router = APIRouter()


# ── Request models ───────────────────────────────────────────

class DiscoveryRunRequest(BaseModel):
    max_count: int = Field(default=25, ge=1, le=500)
    platforms: Optional[list[str]] = None


# ── List discovered targets (date-wise) ─────────────────────

@router.get("")
async def list_discovered_targets(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    run_date: date | None = None,
    platform: str | None = None,
    region: str | None = None,
    is_converted: bool | None = None,
    session: AsyncSession = Depends(get_session),
    _user: User = Depends(get_current_user),
) -> dict:
    """List discovered targets with date-wise filtering and pagination."""
    query = select(DiscoveredTarget)
    count_query = select(func.count()).select_from(DiscoveredTarget)

    if run_date is not None:
        query = query.where(DiscoveredTarget.discovered_on == run_date)
        count_query = count_query.where(DiscoveredTarget.discovered_on == run_date)
    if platform:
        query = query.where(DiscoveredTarget.platform == platform)
        count_query = count_query.where(DiscoveredTarget.platform == platform)
    if region:
        query = query.where(DiscoveredTarget.inferred_region.ilike(f"%{region}%"))
        count_query = count_query.where(DiscoveredTarget.inferred_region.ilike(f"%{region}%"))
    if is_converted is not None:
        query = query.where(DiscoveredTarget.is_converted == is_converted)
        count_query = count_query.where(DiscoveredTarget.is_converted == is_converted)

    total = await session.scalar(count_query) or 0
    offset = (page - 1) * page_size
    result = await session.execute(
        query.order_by(DiscoveredTarget.discovered_on.desc(), DiscoveredTarget.follower_count.desc())
        .offset(offset)
        .limit(page_size)
    )
    items = []
    for dt in result.scalars().all():
        items.append({
            "id": str(dt.id),
            "platform": dt.platform.value,
            "username": dt.username,
            "display_name": dt.display_name,
            "profile_url": dt.profile_url,
            "avatar_url": dt.avatar_url,
            "bio": dt.bio,
            "follower_count": dt.follower_count,
            "is_verified": dt.is_verified,
            "gender": dt.gender,
            "inferred_region": dt.inferred_region,
            "region_confidence": dt.region_confidence,
            "discovered_on": dt.discovered_on.isoformat(),
            "is_converted": dt.is_converted,
            "is_dismissed": dt.is_dismissed,
            "created_at": dt.created_at.isoformat() if dt.created_at else None,
        })

    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": max(1, math.ceil(total / page_size)) if total else 1,
    }


# ── List discovery runs ──────────────────────────────────────

@router.get("/runs")
async def list_discovery_runs(
    limit: int = Query(20, ge=1, le=100),
    session: AsyncSession = Depends(get_session),
    _user: User = Depends(get_current_user),
) -> list[dict]:
    """List discovery run history."""
    result = await session.execute(
        select(DiscoveryRun).order_by(DiscoveryRun.started_at.desc()).limit(limit)
    )
    items = []
    for run in result.scalars().all():
        items.append({
            "id": str(run.id),
            "run_date": run.run_date.isoformat(),
            "platforms": run.platforms or [],
            "total_found": run.total_found,
            "total_converted": run.total_converted,
            "status": run.status,
            "error_message": run.error_message,
            "triggered_by": run.triggered_by,
            "started_at": run.started_at.isoformat() if run.started_at else None,
            "completed_at": run.completed_at.isoformat() if run.completed_at else None,
        })
    return items


# ── Trigger discovery manually ───────────────────────────────

async def _run_discovery_background(max_count: int) -> None:
    """Run discovery in a background task with its own database session."""
    from database.connection import async_session_factory

    if async_session_factory is None:
        discovery_state.emit_log("Database not initialized", level="error")
        discovery_state.reset()
        return

    async with async_session_factory() as session:
        await run_discovery(session, triggered_by="manual", max_count=max_count)


@router.post("/run")
async def trigger_discovery(
    body: DiscoveryRunRequest = DiscoveryRunRequest(),
    _user: User = Depends(require_admin),
) -> dict:
    """Manually trigger a discovery run (returns immediately)."""
    if discovery_state.status in ("running", "paused"):
        raise HTTPException(status_code=409, detail="Discovery is already running")
    discovery_state.mark_starting(body.max_count)
    asyncio.create_task(_run_discovery_background(body.max_count))
    return {"status": "started", "message": f"Discovery started (max: {body.max_count} profiles)"}


# ── Discovery control endpoints ──────────────────────────────

@router.post("/pause")
async def pause_discovery(
    _user: User = Depends(require_admin),
) -> dict:
    """Pause a running discovery."""
    discovery_state.pause()
    return {"status": "paused"}


@router.post("/resume")
async def resume_discovery(
    _user: User = Depends(require_admin),
) -> dict:
    """Resume a paused discovery."""
    discovery_state.resume()
    return {"status": "resumed"}


@router.post("/stop")
async def stop_discovery(
    _user: User = Depends(require_admin),
) -> dict:
    """Stop the running discovery."""
    discovery_state.stop()
    return {"status": "stopped"}


@router.get("/status")
async def get_discovery_status(
    _user: User = Depends(get_current_user),
) -> dict:
    """Return current discovery state."""
    return discovery_state.get_status()


@router.get("/logs/stream")
async def stream_discovery_logs(
    request: Request,
    _user: User = Depends(get_current_user),
):
    """SSE endpoint for streaming discovery logs."""
    async def event_generator():
        last_index = 0
        # Replay existing logs
        logs = discovery_state.get_logs(since_index=0)
        for log in logs:
            yield f"data: {json.dumps(log)}\n\n"
            last_index = log["index"]
        # Stream new logs
        while True:
            if await request.is_disconnected():
                break
            await asyncio.sleep(0.5)
            new_logs = discovery_state.get_logs(since_index=last_index)
            for log in new_logs:
                yield f"data: {json.dumps(log)}\n\n"
                last_index = log["index"]
            # Stop streaming when discovery is idle/stopped and no more logs
            if discovery_state.status in ("idle", "stopped") and not new_logs:
                yield f"data: {json.dumps({'type': 'close', 'message': 'Stream ended'})}\n\n"
                break

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


# ── Convert discovered target to monitored target ────────────

@router.post("/{discovered_id}/convert")
async def convert_to_target(
    discovered_id: UUID,
    session: AsyncSession = Depends(get_session),
    current_user: User = Depends(require_analyst),
) -> dict:
    """Convert a discovered profile into a monitored target account."""
    try:
        target = await convert_discovered_to_target(session, discovered_id, current_user.id)
        return {
            "status": "converted",
            "target_id": str(target.id),
            "username": target.username,
            "platform": target.platform.value,
        }
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc))


# ── Dismiss discovered target ────────────────────────────────

@router.post("/{discovered_id}/dismiss")
async def dismiss_discovered(
    discovered_id: UUID,
    reason: str = Query("", description="Optional reason for dismissal"),
    session: AsyncSession = Depends(get_session),
    _user: User = Depends(require_analyst),
) -> dict:
    """Mark a discovered profile as dismissed (not relevant)."""
    dt = await session.get(DiscoveredTarget, discovered_id)
    if not dt:
        raise HTTPException(status_code=404, detail="Discovered target not found")
    dt.is_dismissed = True
    dt.dismissed_reason = reason or None
    await session.commit()
    return {"status": "dismissed", "id": str(discovered_id)}


# ── Export discovered targets as Excel ───────────────────────

@router.get("/export.xlsx")
async def export_discovered_excel(
    run_date: date | None = None,
    session: AsyncSession = Depends(get_session),
    _user: User = Depends(require_analyst),
) -> StreamingResponse:
    """Export discovered targets as Excel file."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        # Fallback to CSV if openpyxl not installed
        return await _export_discovered_csv(run_date, session)

    query = select(DiscoveredTarget)
    if run_date:
        query = query.where(DiscoveredTarget.discovered_on == run_date)
    query = query.order_by(DiscoveredTarget.discovered_on.desc(), DiscoveredTarget.follower_count.desc())
    result = await session.execute(query.limit(5000))

    wb = Workbook()
    ws = wb.active
    ws.title = "Discovered Targets"

    # Header
    headers = [
        "Date", "Platform", "Username", "Display Name", "Profile URL",
        "Followers", "Verified", "Gender", "Region", "Region Confidence",
        "Bio", "Converted", "Status",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, dt in enumerate(result.scalars().all(), 2):
        ws.cell(row=row_idx, column=1, value=dt.discovered_on.isoformat())
        ws.cell(row=row_idx, column=2, value=dt.platform.value)
        ws.cell(row=row_idx, column=3, value=dt.username)
        ws.cell(row=row_idx, column=4, value=dt.display_name or "")
        ws.cell(row=row_idx, column=5, value=dt.profile_url)
        ws.cell(row=row_idx, column=6, value=dt.follower_count)
        ws.cell(row=row_idx, column=7, value="Yes" if dt.is_verified else "No")
        ws.cell(row=row_idx, column=8, value=dt.gender)
        ws.cell(row=row_idx, column=9, value=dt.inferred_region or "")
        ws.cell(row=row_idx, column=10, value=round(dt.region_confidence or 0, 2))
        ws.cell(row=row_idx, column=11, value=(dt.bio or "")[:200])
        ws.cell(row=row_idx, column=12, value="Yes" if dt.is_converted else "No")
        ws.cell(row=row_idx, column=13, value="Converted" if dt.is_converted else ("Dismissed" if dt.is_dismissed else "Pending"))

    # Auto-width columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    date_label = run_date.isoformat() if run_date else "all"
    filename = f"shakti_discovered_targets_{date_label}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export.csv")
async def export_discovered_csv_route(
    run_date: date | None = None,
    session: AsyncSession = Depends(get_session),
    _user: User = Depends(require_analyst),
) -> StreamingResponse:
    """Export discovered targets as CSV file."""
    return await _export_discovered_csv(run_date, session)


async def _export_discovered_csv(
    run_date: date | None,
    session: AsyncSession,
) -> StreamingResponse:
    """Internal CSV export helper."""
    query = select(DiscoveredTarget)
    if run_date:
        query = query.where(DiscoveredTarget.discovered_on == run_date)
    query = query.order_by(DiscoveredTarget.discovered_on.desc(), DiscoveredTarget.follower_count.desc())
    result = await session.execute(query.limit(5000))

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([
        "Date", "Platform", "Username", "Display Name", "Profile URL",
        "Followers", "Verified", "Gender", "Region", "Bio", "Converted",
    ])
    for dt in result.scalars().all():
        writer.writerow([
            dt.discovered_on.isoformat(),
            dt.platform.value,
            dt.username,
            dt.display_name or "",
            dt.profile_url,
            dt.follower_count,
            "Yes" if dt.is_verified else "No",
            dt.gender,
            dt.inferred_region or "",
            (dt.bio or "")[:200],
            "Yes" if dt.is_converted else "No",
        ])

    content = buffer.getvalue().encode("utf-8-sig")
    date_label = run_date.isoformat() if run_date else "all"
    filename = f"shakti_discovered_targets_{date_label}.csv"
    return StreamingResponse(
        io.BytesIO(content),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
